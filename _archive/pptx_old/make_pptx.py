"""
White PPTX decks — clean layout, no mashed clipart.
User AI images: put files in assets/user/ then re-run this script.
  logo.png, temple.png, guide.png, ritual_thumb.png, music_cover.png
Live screenshots (optional): assets/screenshots/*.png
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DIR = Path(__file__).resolve().parent
ASSETS = DIR / "assets"
USER = ASSETS / "user"
SHOTS = ASSETS / "screenshots"

BG = RGBColor(255, 255, 255)
INK = RGBColor(26, 26, 26)
MUTED = RGBColor(90, 90, 90)
LINE = RGBColor(220, 220, 220)
SOFT = RGBColor(248, 248, 248)
ACCENT = RGBColor(44, 95, 110)
WIP = RGBColor(160, 100, 40)


def font(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def set_bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    return s


def label(slide, text, top=0.45, color=ACCENT):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.5), Inches(0.35))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    font(r, 13, True, color)


def title(slide, text, top=0.9, size=32, color=INK):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.5), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = text
    font(r, size, True, color)


def body(slide, lines, top=2.2, size=18, color=MUTED, left=0.85, width=11.5):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(4.4))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = line
        font(r, size, False, color)


def footer(slide, name, n, total):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(7.05), Inches(11.5), Inches(0.3))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = f"{name}  ·  인간과 종교  ·  {n}/{total}"
    font(r, 11, False, ACCENT)


def card(slide, left, top, width, height, heading, lines):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT
    shape.line.color.rgb = LINE
    box = slide.shapes.add_textbox(
        Inches(left + 0.22), Inches(top + 0.22), Inches(width - 0.4), Inches(height - 0.4)
    )
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = heading
    font(r, 16, True, ACCENT)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = line
        font(r, 14, False, MUTED)


def placeholder(slide, left, top, width, height, caption):
    """Empty slot for user-provided AI image."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.adjustments[0] = 0.04
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT
    shape.line.color.rgb = LINE
    box = slide.shapes.add_textbox(Inches(left), Inches(top + height / 2 - 0.35), Inches(width), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = caption
    font(r, 14, False, MUTED)


def one_image(slide, path, left, top, width, height, fallback_caption):
    p = Path(path) if path else None
    if p and p.exists():
        slide.shapes.add_picture(str(p), Inches(left), Inches(top), width=Inches(width))
    else:
        placeholder(slide, left, top, width, height, fallback_caption)


def make_plan():
    prs = new_prs()
    total = 6
    name = "계획안"

    s = add(prs)
    label(s, "인간과 종교  ·  조별 계획안  ·  9월 17일")
    title(s, "공명(共鳴) 스테이션", size=40)
    body(
        s,
        [
            "2035 미래 종교 디자인 — 계획안",
            "「말해도 된다. 여기는 들어준다.」",
            "",
            "자세한 본문: 계획안_공명스테이션.md",
        ],
        top=2.5,
        size=20,
        color=INK,
    )
    one_image(s, USER / "logo.png", 9.5, 2.3, 3.0, 3.0, "AI 이미지 자리\n로고 → assets/user/logo.png")
    footer(s, name, 1, total)

    s = add(prs)
    label(s, "특성 · 이름 · 세계관")
    title(s, "무엇을 디자인하는가", size=30)
    card(s, 0.85, 2.2, 3.8, 4.0, "미래 종교 특성", ["① 돌봄 (Micro-Care)", "② 일상의 의례", "③ AI는 매개(Bridge)"])
    card(s, 4.85, 2.2, 3.8, 4.0, "이름", ["공명 스테이션", "새 교단 창시 아님", "다종교 영적 인프라"])
    card(s, 8.85, 2.2, 3.6, 4.0, "세계관", ["신 = 연대의 순간", "거룩함 =", "평가 없는 취약함의 발화"])
    footer(s, name, 2, total)

    s = add(prs)
    label(s, "상징 · 의례 · 공간")
    title(s, "말이 울리는 자리", size=30)
    body(
        s,
        [
            "상징 — 파동(공명) · 빈 의자(받아들여짐)",
            "의례 — 내려놓음의 대화: 앉기 → 말하기 → 침묵 → 위로 → 연대 → 휘발",
            "공간 — 복지관·구청 로비의 1인 방음 부스 (= 성전)",
        ],
        top=2.15,
        size=18,
        color=INK,
        width=7.2,
    )
    one_image(s, USER / "temple.png", 8.5, 2.1, 4.0, 4.2, "AI 이미지 자리\n성전/부스\nassets/user/temple.png")
    footer(s, name, 3, total)

    s = add(prs)
    label(s, "공동체 · 윤리 · AI")
    title(s, "어떻게 돌보는가", size=30)
    card(s, 0.85, 2.2, 5.7, 4.0, "공동체 · 윤리", ["익명·느슨한 연대", "평가 금지 · 원문 휘발", "위험 시에만 인간 연결"])
    card(s, 6.8, 2.2, 5.7, 4.0, "AI 역할", ["1차 경청 · 짧은 위로", "신은 되지 않음", "Bridge → 성직자·복지사"])
    footer(s, name, 4, total)

    s = add(prs)
    label(s, "계획 · 예상 결과")
    title(s, "앞으로 만들 것", size=30)
    body(
        s,
        [
            "AI 활용: 로고 · 성전 · 의례 영상 · 음악 · 가상 지도자 · 웹 프로토타입",
            "",
            "10/1 중간 — 제작 중인 사이트에서 의례 일부 시연",
            "10/8 최종 — 포함 요소·AI 산출물·완성된 시연",
        ],
        top=2.2,
        size=19,
        color=INK,
    )
    footer(s, name, 5, total)

    s = add(prs)
    label(s, "닫기")
    title(s, "혼자 앓던 말이, 여기서 울린다.", top=2.8, size=34, color=ACCENT)
    body(s, ["질문 환영합니다."], top=4.3, size=20)
    footer(s, name, 6, total)

    out = DIR / "계획안_0917.pptx"
    prs.save(out)
    print("wrote", out)


def make_mid():
    """In-progress midterm — not a finished product pitch."""
    prs = new_prs()
    total = 6
    name = "중간 · 제작 중"

    s = add(prs)
    label(s, "인간과 종교  ·  중간 결과물  ·  10월 1일", color=WIP)
    title(s, "지금은 ‘완성’이 아니라 ‘제작 중’입니다", size=30)
    body(
        s,
        [
            "중간 목표: 의례의 핵심(경청 → 위로)이 웹에서 돌아가는지 확인",
            "아직 없는 것: 로고·성전 시각 확정, 의례 영상·음악, 공간 연출 완성",
        ],
        top=2.5,
        size=20,
        color=INK,
    )
    footer(s, name, 1, total)

    s = add(prs)
    label(s, "계획 대비 진행", color=WIP)
    title(s, "어디까지 왔는가", size=30)
    card(
        s,
        0.85,
        2.2,
        5.7,
        4.0,
        "된 것 (프로토타입)",
        [
            "이름·세계관 초안 확정",
            "내려놓음 대화 흐름 구현 중",
            "전통 선택 · 음성/텍스트 경청",
            "짧은 위로 응답 실험",
        ],
    )
    card(
        s,
        6.8,
        2.2,
        5.7,
        4.0,
        "아직 제작 중",
        [
            "상징·로고 AI 확정",
            "성전(부스) 시각·영상",
            "익명 연대 완성도",
            "음악 · 가상 지도자 페르소나",
        ],
    )
    footer(s, name, 2, total)

    s = add(prs)
    label(s, "중간 시연 — 화면 1장만", color=WIP)
    title(s, "키오스크 진입 (제작 중인 UI)", size=28)
    # Single screenshot only — no mashup
    one_image(
        s,
        SHOTS / "00_home.png",
        2.5,
        2.0,
        8.3,
        4.6,
        "스크린샷 자리\nassets/screenshots/00_home.png",
    )
    footer(s, name, 3, total)

    s = add(prs)
    label(s, "중간 시연 — 의례 핵심", color=WIP)
    title(s, "경청 → 위로가 이어지는 장면", size=28)
    one_image(
        s,
        SHOTS / "05_chat.png",
        2.5,
        2.0,
        8.3,
        4.6,
        "스크린샷 자리\nassets/screenshots/05_chat.png",
    )
    footer(s, name, 4, total)

    s = add(prs)
    label(s, "중간에서 확인한 점", color=WIP)
    title(s, "종교적 방식은 살아 있다", size=28)
    body(
        s,
        [
            "해결·설교보다 반영과 수용이 가능하다",
            "말이 끊겨도 ‘들어주는’ 의례로 설계 중",
            "다만 UI·상징·영상은 최종에서 다듬는다",
        ],
        top=2.4,
        size=20,
        color=INK,
    )
    footer(s, name, 5, total)

    s = add(prs)
    label(s, "최종을 향해", color=WIP)
    title(s, "10/8에 채울 것", size=30)
    body(
        s,
        [
            "① AI 로고·성전·지도자·음악·의례 영상",
            "② 사이트 시연 완성도",
            "③ 과제 포함 요소 한눈에 매핑",
            "",
            "오늘은 ‘만드는 과정’을 보여 드렸습니다.",
        ],
        top=2.3,
        size=20,
        color=INK,
    )
    footer(s, name, 6, total)

    out = DIR / "중간결과_1001.pptx"
    prs.save(out)
    print("wrote", out)


def make_final():
    """Fully developed final — distinct from mid WIP."""
    prs = new_prs()
    total = 9
    name = "최종"

    s = add(prs)
    label(s, "인간과 종교  ·  최종 결과물  ·  10월 8일")
    title(s, "공명 스테이션 — 완성된 공동체 디자인", size=30)
    body(
        s,
        [
            "2035 영적 인프라 · 포함 요소 · AI 활용 · 실제 시연",
            "「말해도 된다. 여기는 들어준다.」",
        ],
        top=2.6,
        size=20,
        color=INK,
    )
    one_image(s, USER / "logo.png", 9.6, 2.4, 2.9, 2.9, "로고\nassets/user/logo.png")
    footer(s, name, 1, total)

    s = add(prs)
    label(s, "이름 · 세계관")
    title(s, "새 교단이 아니라, 돌봄의 인프라", size=28)
    body(
        s,
        [
            "이름: 공명(共鳴) 스테이션",
            "세계관: Micro-Care · 신은 연대의 순간 · AI는 Bridge",
            "거룩함: 평가 없이 취약함을 소리 내는 행위",
            "기존 종교와의 관계: 대체재가 아니라 연장재",
        ],
        top=2.3,
        size=19,
        color=INK,
        width=7.5,
    )
    one_image(s, USER / "logo.png", 8.8, 2.3, 3.6, 3.6, "상징/로고")
    footer(s, name, 2, total)

    s = add(prs)
    label(s, "상징 · 성스러운 공간")
    title(s, "파동 · 빈 의자 · 1인 부스", size=28)
    body(
        s,
        [
            "상징: 말이 울림(파동), 받아들여질 자리(빈 의자)",
            "공간: 복지관·구청·병원·종교시설 부속의 방음 부스",
            "성스러움: 화려한 성물보다 침묵·차단·느린 빛",
        ],
        top=2.2,
        size=18,
        color=INK,
        width=6.5,
    )
    one_image(s, USER / "temple.png", 7.8, 2.1, 4.7, 4.3, "성전/부스\nassets/user/temple.png")
    footer(s, name, 3, total)

    s = add(prs)
    label(s, "의례")
    title(s, "내려놓음의 대화", size=30)
    body(
        s,
        [
            "앉기 → 전통 선택 → 말하기 → 「들었습니다」와 침묵",
            "→ 짧은 위로 → 익명 연대(선택) → 휘발",
            "",
            "해결·개종 권유 없음. 수용이 의례의 중심.",
        ],
        top=2.2,
        size=19,
        color=INK,
        width=6.8,
    )
    one_image(s, USER / "ritual_thumb.png", 8.0, 2.1, 4.5, 4.3, "의례 영상 썸네일\nassets/user/ritual_thumb.png")
    footer(s, name, 4, total)

    s = add(prs)
    label(s, "공동체 · 윤리 · AI")
    title(s, "익명 연대 · Bridge", size=28)
    card(s, 0.85, 2.2, 5.7, 4.0, "공동체 · 윤리", ["익명 연대 — 나만 그런 게 아님", "평가 금지 · 원문 휘발", "위험 시 인간에게 핸드오프"])
    one_image(s, USER / "guide.png", 7.0, 2.1, 5.4, 4.2, "가상 종교지도자(Bridge)\nassets/user/guide.png")
    footer(s, name, 5, total)

    s = add(prs)
    label(s, "AI 활용 결과물")
    title(s, "과제에서 요구한 시각·청각 산출물", size=26)
    # Five separate slots — not mashed into one collage
    one_image(s, USER / "logo.png", 0.7, 2.2, 2.3, 2.3, "로고")
    one_image(s, USER / "temple.png", 3.2, 2.2, 2.3, 2.3, "성전")
    one_image(s, USER / "ritual_thumb.png", 5.7, 2.2, 2.3, 2.3, "의례 영상")
    one_image(s, USER / "music_cover.png", 8.2, 2.2, 2.3, 2.3, "음악")
    one_image(s, USER / "guide.png", 10.7, 2.2, 2.0, 2.3, "지도자")
    body(s, ["파일을 assets/user/ 에 넣고 make_pptx.py 를 다시 실행하면 자동 삽입됩니다."], top=5.0, size=14)
    footer(s, name, 6, total)

    s = add(prs)
    label(s, "실제 시연")
    title(s, "사이트 작동 — 한 장면씩", size=28)
    one_image(s, SHOTS / "00_home.png", 0.7, 2.0, 5.9, 4.4, "홈 스크린샷")
    one_image(s, SHOTS / "05_chat.png", 6.9, 2.0, 5.7, 4.4, "위로 대화 스크린샷")
    footer(s, name, 7, total)

    s = add(prs)
    label(s, "과제 매핑 · 평가")
    title(s, "포함 요소와 평가지표", size=28)
    body(
        s,
        [
            "포함: 이름·세계관·상징·의례·공동체·윤리·공간·AI역할",
            "AI: 로고·성전·의례영상·음악·가상지도자·웹",
            "",
            "설득력 — 거룩함·의례·연장재   |   창의 — Bridge·시각물",
            "활용성 — 일상 부스·접근성   |   완성도 — 문서·PPT·시연",
        ],
        top=2.3,
        size=18,
        color=INK,
    )
    footer(s, name, 8, total)

    s = add(prs)
    label(s, "닫기")
    title(s, "혼자 앓던 말이, 여기서 울린다.", top=2.8, size=34, color=ACCENT)
    body(s, ["질문 환영합니다."], top=4.3, size=20)
    footer(s, name, 9, total)

    out = DIR / "최종결과_1008.pptx"
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    USER.mkdir(parents=True, exist_ok=True)
    make_plan()
    make_mid()
    make_final()
