# -*- coding: utf-8 -*-
"""계획안 PPTX — HTML과 동일 14장·상세 8요소."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "발표_미래종교와_공명.pptx"

BG = RGBColor(243, 244, 246)
SURFACE = RGBColor(255, 255, 255)
INK = RGBColor(20, 24, 32)
MUTED = RGBColor(75, 85, 99)
ACCENT = RGBColor(15, 118, 110)
LINE = RGBColor(220, 224, 230)
SOFT = RGBColor(215, 239, 236)


def font(run, size=16, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG


def add(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


def tag(s, text):
    b = s.shapes.add_textbox(Inches(0.65), Inches(0.32), Inches(12), Inches(0.28))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = text
    font(r, 11, True, ACCENT)


def title(s, text, top=0.58, size=30):
    b = s.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.1), Inches(0.75))
    tf = b.text_frame
    tf.word_wrap = True
    first = True
    for line in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line
        font(r, size, True, INK)


def explain(s, text, top=1.28):
    b = s.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.1), Inches(0.55))
    tf = b.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = text
    font(r, 16, False, MUTED)


def footer(s, n, total, label=""):
    b = s.shapes.add_textbox(Inches(0.65), Inches(7.08), Inches(12), Inches(0.28))
    r = b.text_frame.paragraphs[0].add_run()
    r.text = f"{label}  ·  {n} / {total}" if label else f"조별 계획안 · 공명 스테이션  ·  {n} / {total}"
    font(r, 11, False, MUTED)


def card(s, left, top, w, h, head, lines, head_size=18, body_size=15):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = SURFACE
    sh.line.color.rgb = LINE
    box = s.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.18), Inches(w - 0.4), Inches(h - 0.36))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = head
    font(r, head_size, True, INK)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = line
        font(r, body_size, False, MUTED)


def kpi(s, left, top, w, h, label, value, hint):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = SURFACE
    sh.line.color.rgb = LINE
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.07), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    box = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(w - 0.45), Inches(h - 0.4))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = label
    font(r, 12, True, ACCENT)
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run()
    r.text = value
    font(r, 22, True, INK)
    p = tf.add_paragraph()
    p.space_before = Pt(8)
    r = p.add_run()
    r.text = hint
    font(r, 14, False, MUTED)


def quote_bar(s, top, text):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(top), Inches(12.05), Inches(0.7))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = SURFACE
    sh.line.color.rgb = LINE
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(top), Inches(0.08), Inches(0.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.95), Inches(top + 0.15), Inches(11.5), Inches(0.45))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    font(r, 16, True, INK)


def body_lines(s, lines, top, size=17, color=INK, h=4.5):
    b = s.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.1), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        font(r, size, False, color)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    T = 14

    # 1
    s = add(prs)
    tag(s, "PLAN · FUTURE RELIGION DESIGN")
    title(s, "들어줄 자리가 줄었다.\n그래서 인프라를 깐다.", top=1.8, size=32)
    body_lines(s, [
        "공명 스테이션 — 2035, 새 교단이 아니라 다종교가 공유하는 영적 인프라.",
        "이름 · 세계관 · 상징 · 의례 · 공동체 · 윤리 · 공간 · AI",
    ], top=3.6, size=16, color=MUTED, h=2)
    footer(s, 1, T, "계획안")

    # 2
    s = add(prs)
    tag(s, "AGENDA")
    title(s, "오늘 이야기 — 8요소가 중심")
    explain(s, "과제 조건·문제 의식은 짧게, 종교 형태 8요소를 자세히 설명한다. 결과물 형태는 끝에서 한 장만.")
    body_lines(s, [
        "01  이름 — 공명 스테이션 · 뜻·성격·표어",
        "02  세계관 — Micro-Care · 신 · 다종교 · 거룩함",
        "03  상징 — 파동 · 빈 의자 · 로고",
        "04  의례 — 내려놓음의 대화 10단계",
        "05  공동체 — 익명 연대 · 기존 종교와 관계",
        "06  윤리 — 평가 금지 · 휘발 · 위험 · 규모",
        "07  성스러운 공간 — 1인 방음 부스",
        "08  AI 역할 — Bridge (경청·연결·매개)",
    ], top=1.85, size=16, color=INK, h=4.8)
    footer(s, 2, T, "목차")

    # 3
    s = add(prs)
    tag(s, "ASSIGNMENT")
    title(s, "과제가 요구하는 것")
    soft = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.35), Inches(12.05), Inches(0.95))
    soft.adjustments[0] = 0.06
    soft.fill.solid()
    soft.fill.fore_color.rgb = SOFT
    soft.line.color.rgb = ACCENT
    box = s.shapes.add_textbox(Inches(0.85), Inches(1.5), Inches(11.6), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = "미래종교 디자인 — 2035, 종교·문화·기술이 결합한 공동체"
    font(r, 16, True, INK)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "핵심은 종교 형태 설계(8요소). 결과물은 그걸 보여주는 수단."
    font(r, 13, False, MUTED)
    body_lines(s, [
        "포함 8  이름 · 세계관 · 상징 · 의례 · 공동체 · 윤리 · 성스러운 공간 · AI 역할",
        "AI 활용  로고 · 성전 · 의례 영상 · 음악 · 가상 종교지도자 등",
        "제출  결과물 최소 2개 (웹·챗봇·영상·음악·게임 등)",
        "평가  종교학적 설득력 · AI 창의성 · 현실 활용성 · 완성도 · 팀웍",
        "본 조  공명 스테이션 — 새 교단이 아닌, 다종교·지역사회용 영적 인프라",
    ], top=2.5, size=15, color=INK, h=4)
    footer(s, 3, T, "교과 안내")

    # 4
    s = add(prs)
    tag(s, "WHY")
    title(s, "왜 이런 종교 형태인가")
    explain(s, "과학이 설명을 맡아가도 의미·애도·경청·소속의 필요는 남는다. 연결은 늘었지만 평가 없이 끝까지 듣는 자리는 줄었다.")
    kpi(s, 0.65, 1.85, 3.95, 4.2, "설명 ≠ 의미", "의미는 남음",
        "세계·몸·규범의 설명은 과학·의학·법률이 가져갔다. 그러나 애도하고, 듣고, 소속될 필요는 사라지지 않았다.")
    kpi(s, 4.75, 1.85, 3.95, 4.2, "연결 ≠ 경청", "들어줄 자리 ↓",
        "기술은 사람을 연결하지만, 가르치지 않고 끝까지 들어 주는 자리는 희소하다. 고령·독거에서 특히 드러난다.")
    kpi(s, 8.85, 1.85, 3.85, 4.2, "설계 방향", "Micro-Care",
        "사실을 이기는 장치가 아니라, 취약함을 평가 없이 들어 주는 돌봄의 인프라로 종교 형태를 옮긴다.")
    footer(s, 4, T, "문제 의식")

    # 5 이름
    s = add(prs)
    tag(s, "01 · NAME")
    title(s, "이름 — 공명(共鳴) 스테이션")
    explain(s, "특정 종교의 브랜드가 아니라, 여러 전통이 함께 쓸 수 있는 인프라의 이름이다.")
    card(s, 0.65, 1.85, 6.05, 4.6, "이름이 뜻하는 것", [
        "공명(共鳴) — 말이 평가·교정·개종 권유 없이 받아들여지는 상태.",
        "“잘 울린다”는 감성이 아니라, 경청이 성립했다는 뜻.",
        "",
        "스테이션 — 그 일이 일상적으로 일어나는 자리.",
        "예배당에만 가지 않아도 내려놓을 수 있는 거점.",
        "",
        "표어: 「듣고 있습니다」 · 「말씀해 주십시오」",
    ], body_size=15)
    card(s, 6.9, 1.85, 5.8, 4.6, "이 이름의 성격", [
        "새 교단·신흥 종교를 창시하지 않는다.",
        "가톨릭·개신교·불교·원불교·무종교 등이 공유.",
        "기존 종교의 대체재가 아니라 연장재.",
        "못 가는 날, 말이 막히는 날의 Micro-Care.",
        "부스 안: 교리 시험·개종 권유 없음.",
        "소속은 강요하지 않는다.",
    ], body_size=15)
    footer(s, 5, T, "이름")

    # 6 세계관
    s = add(prs)
    tag(s, "02 · WORLDVIEW")
    title(s, "세계관 — 무게중심이 이동한다")
    explain(s, "2035 초고령·고립. 신은 먼 절대자만이 아니라, 취약함이 받아들여지는 연대의 순간에 있다.")
    card(s, 0.65, 1.85, 5.5, 3.3, "이전의 무게중심", [
        "거대한 진리·교리의 전파",
        "예배당·사찰에 모여야 성스러움",
        "신은 저 너머의 절대자",
        "AI = 사제·신을 대체",
    ], body_size=15)
    card(s, 7.15, 1.85, 5.55, 3.3, "공명의 무게중심", [
        "Micro-Care — 미시적 경청과 위로",
        "일상 부스에서 내려놓음이 의례",
        "신은 연대의 순간에 현존",
        "AI = Bridge (매개, 대체 아님)",
    ], body_size=15)
    card(s, 0.65, 5.3, 6.05, 1.4, "다종교 구조", [
        "공통 레이어(공간·윤리·휘발·위험 연결)는 같고,",
        "전통 레이어는 위로 어휘만 다르다. 교리 시험·개종 없음.",
    ], body_size=16)
    card(s, 6.9, 5.3, 5.8, 1.4, "거룩함의 정의", [
        "기도문 암송이 아니라,",
        "평가 없이 취약함을 소리 내는 행위 자체가 거룩한 의례.",
    ], body_size=16)
    footer(s, 6, T, "세계관")

    # 7 상징
    s = add(prs)
    tag(s, "03 · SYMBOL")
    title(s, "상징 — 소속 표시가 아니라 수용의 표시")
    explain(s, "특정 종교의 성물로 “우리 편”을 드러내지 않는다. 말이 받아들여지고 있다는 감각만 남긴다.")
    kpi(s, 0.65, 1.85, 3.95, 4.0, "상징 01", "파동 (Wave)",
        "말할 때 부드럽게 반응, 침묵 때 느린 맥동. “지금 듣고 있다”는 시각 피드백. 소리가 허공에 떨어지지 않았다는 표시.")
    kpi(s, 4.75, 1.85, 3.95, 4.0, "상징 02", "빈 의자",
        "십자가·불상 등 특정 성물 비배치. 앉는 행위 자체가 “받아들여질 자리에 왔다”는 상징. 누구나 올 수 있는 빈자리.")
    kpi(s, 8.85, 1.85, 3.85, 4.0, "상징 03", "로고",
        "파동과 여백 중심 마크. AI 제작으로 과제 AI 활용(상징물)에 대응. 화려함보다 비움·반응이 보이게.")
    quote_bar(s, 6.05, "상징이 교단을 광고하지 않는다. 상징은 “여기서는 평가하지 않는다”를 보여 준다.")
    footer(s, 7, T, "상징")

    # 8 의례
    s = add(prs)
    tag(s, "04 · RITUAL")
    title(s, "의례 — 「내려놓음의 대화」 · 약 5–8분")
    explain(s, "기도문 암송이 아니라 소리 내어 말하기가 거룩한 행위. 핵심은 발화→침묵→반영 위로(④–⑥). *는 선택.")
    steps = [
        ("1 입장", "방음·빈 의자. 앉는 순간이 시작."),
        ("2 전통", "전통/공통 선택. 어휘만 다름."),
        ("3 초대", "「말씀해 주십시오」"),
        ("4 발화", "끊겨도 됨. 평가·진단 금지."),
        ("5 침묵", "듣고 있습니다. 여백."),
        ("6 위로", "반영 + 전통 짧은 문장."),
        ("7 연대*", "익명 짧은 목소리."),
        ("8 실천*", "카드 1장. 스킵 OK."),
        ("9 사람*", "Bridge — 다음 문."),
        ("10 휘발", "일반 대화 원문 삭제."),
    ]
    for i, (h, d) in enumerate(steps):
        col = i % 5
        row = i // 5
        left = 0.65 + col * 2.5
        top = 1.85 + row * 2.35
        card(s, left, top, 2.35, 2.15, h, [d], head_size=16, body_size=14)
    footer(s, 8, T, "의례")

    # 9 공동체
    s = add(prs)
    tag(s, "05 · COMMUNITY")
    title(s, "공동체 — 익명의 느슨한 연대")
    explain(s, "정기 집회·교단 가입이 아니다. “나만 그런 게 아니다”를 느끼는 감각이 공동체다.")
    card(s, 0.65, 1.85, 3.95, 4.6, "기본 형태", [
        "이름·얼굴·연락처 교환 없음이 기본",
        "동의 시에만 익명 클립",
        "(성문 제거·기한부·태그 1개)",
        "비슷한 밤을 보낸 이의 말을 짧게",
        "출석·전화·세력화로 모으지 않음",
    ], body_size=15)
    card(s, 4.75, 1.85, 3.95, 4.6, "왜 느슨한가", [
        "초고령·거동·수치심으로",
        "물리 모임이 어려운 사람을 염두",
        "강제 소속은 다시 문턱이 된다",
        "고립을 깨는 최소 단위",
        "= “혼자가 아님”의 감각",
    ], body_size=15)
    card(s, 8.85, 1.85, 3.85, 4.6, "기존 종교와", [
        "대체하지 않는다 — 연장재",
        "(선택) 성당·교회·사찰·복지 안내",
        "깊은 사목·성례는 인간에게",
        "부스 안 개종·교리 논쟁 없음",
    ], body_size=15)
    footer(s, 9, T, "공동체")

    # 10 윤리
    s = add(prs)
    tag(s, "06 · ETHICS")
    title(s, "윤리 — 다섯 원칙 (왜 필요한가)")
    explain(s, "경청만으로는 부족하다. 해악·위험·규모를 막지 않으면 돌봄이 우상·세력·방치가 된다.")
    kpi(s, 0.65, 1.85, 6.05, 2.3, "01 · 평가 금지", "No Judgment",
        "비난·성급한 긍정·개종 권유·확정 진단 금지. 먼저 듣고 상대 말을 반영한다.")
    kpi(s, 6.9, 1.85, 5.8, 2.3, "02 · 휘발", "말하지 않은 것처럼",
        "일반 세션 원문은 남기지 않는다. “밖으로 새지 않는다”는 신뢰가 의례의 조건.")
    kpi(s, 0.65, 4.3, 6.05, 2.3, "03 · 위험 예외", "최소 정보로 사람",
        "자해·타해 등 → 휘발보다 안전. 성직자·복지사·핫라인에 연결.")
    kpi(s, 6.9, 4.3, 5.8, 2.3, "04–05 · 해악 · 규모", "단호함 · 경계",
        "타해·혐오·복수 미화는 위로하지 않음. 모으지 않기 / 세력화·AI 숭배 금지.")
    footer(s, 10, T, "윤리")

    # 11 공간
    s = add(prs)
    tag(s, "07 · SACRED SPACE")
    title(s, "성스러운 공간 — 1인 방음 부스")
    explain(s, "“교회처럼 보이는 건물”이 아니라, 말이 새어 나가지 않는 방이 성스럽다.")
    card(s, 0.65, 1.85, 6.05, 4.6, "어디에 두는가", [
        "노인정 · 복지관 · 구청 로비",
        "병원 로비 · 종교시설 부속 공간",
        "예배당만이 아니라 생활 동선",
        "거동·교통·수치심 때문에 못 가는 날을 염두",
        "한 번에 한 사람. 집회장이 아니다",
    ], body_size=16)
    card(s, 6.9, 1.85, 5.8, 4.6, "무엇이 성스러운가", [
        "침묵과 외부 소음·시선 차단",
        "빈 의자 — 받아들여질 자리",
        "느린 빛 · 과도한 장식·성물 비배치",
        "「말이 안전히 떨어지는 방」",
        "성스러움 = 평가 없이 말할 수 있는 조건",
        "(건축의 화려함이 아님)",
    ], body_size=16)
    footer(s, 11, T, "성스러운 공간")

    # 12 AI
    s = add(prs)
    tag(s, "08 · AI ROLE")
    title(s, "AI 역할 — Bridge")
    explain(s, "AI는 신·사제 대체가 아니라, 1차 경청 후 사람에게 넘기는 매개다.")
    kpi(s, 0.65, 1.85, 3.95, 4.0, "역할 1 · 1차 경청", "반영하고 위로",
        "감정·사건·몸·관계를 되받은 뒤 짧게 위로. 전통 어휘는 실제로 쓰이는 한 문장. 교리 강의·숙제 금지.")
    kpi(s, 4.75, 1.85, 3.95, 4.0, "역할 2 · Bridge", "사람에게 넘김",
        "위험·깊은 사목·원함 → 성직자·복지사·핫라인. 실패가 아니라 의례의 다음 문.")
    kpi(s, 8.85, 1.85, 3.85, 4.0, "역할 3 · 매개", "지도자 ≠ 신",
        "가상 종교지도자는 숭배 대상이 아님. 경청을 돕는 시각적 매개. AI 숭배 금지.")
    quote_bar(s, 6.05, "기술이 신을 대체하면 안 된다. 기술은 들어줄 자리를 넓히고, 책임은 사람에게 남긴다.")
    footer(s, 12, T, "AI 역할")

    # 13 form
    s = add(prs)
    tag(s, "FORM · 한 장만")
    title(s, "어떻게 보여 줄 것인가")
    explain(s, "위 8요소가 본론. 아래는 그걸 체험·제출로 옮기는 형태만.")
    items = [
        ("이름", "공명 스테이션"), ("세계관", "Micro-Care · 연대"),
        ("상징", "파동 · 빈 의자"), ("의례", "내려놓음의 대화"),
        ("공동체", "익명 연대"), ("윤리", "평가금지 · 휘발"),
        ("공간", "1인 방음 부스"), ("AI", "Bridge"),
    ]
    for i, (k, v) in enumerate(items):
        col, row = i % 4, i // 4
        left = 0.65 + col * 3.15
        top = 1.85 + row * 1.55
        card(s, left, top, 3.0, 1.4, k, [v], head_size=14, body_size=16)
    card(s, 0.65, 5.15, 6.05, 1.5, "제출 ①", ["웹/키오스크 — 의례 흐름(경청→위로→연대)을 보여 줌"], body_size=15)
    card(s, 6.9, 5.15, 5.8, 1.5, "제출 ②", ["의례 영상 또는 AI 이미지 (+ 로고·성전·가상 지도자)"], body_size=15)
    footer(s, 13, T, "결과물 형태")

    # 14
    s = add(prs)
    tag(s, "CLOSING")
    title(s, "한 줄", top=1.5, size=28)
    body_lines(s, [
        "들어줄 존재가 부족한 사회에서,",
        "종교의 역할을 Micro-Care 인프라로 옮긴다.",
    ], top=2.3, size=20, color=INK, h=1.4)
    kpi(s, 0.65, 4.0, 3.95, 2.3, "성격", "새 교단 아님", "다종교 연장재")
    kpi(s, 4.75, 4.0, 3.95, 2.3, "핵심", "듣고 있습니다", "평가 없는 경청")
    kpi(s, 8.85, 4.0, 3.85, 2.3, "AI", "Bridge", "신은 되지 않음")
    footer(s, 14, T, "질문 환영")

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
